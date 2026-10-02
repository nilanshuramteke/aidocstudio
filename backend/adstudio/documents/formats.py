"""Native-text paginators for office/data/email formats. Each returns (pages, meta); meta may carry `_attachments`.

Parsers only read: no macros, no formula evaluation (cached values only), XML via defusedxml.
"""
import json
import re

from ..core.errors import NonRetryable

MAX_TEXT_CHARS = 3_000_000  # per page; huge sheets/CSVs are only useful as searchable text
MAX_SHEET_ROWS = 50_000


def _native(text: str) -> dict:
    if len(text) > MAX_TEXT_CHARS:
        text = text[:MAX_TEXT_CHARS] + "\n[truncated]"
    return {"w": None, "h": None, "text": text, "src": "native"}


def paginate_xlsx(path):
    import openpyxl
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:  # noqa: BLE001
        raise NonRetryable(f"Unreadable XLSX: {e}") from e
    pages, names = [], []
    try:
        for ws in wb.worksheets:
            lines = []
            for i, row in enumerate(ws.iter_rows(values_only=True)):
                if i >= MAX_SHEET_ROWS:
                    lines.append("[more rows not indexed]")
                    break
                cells = ["" if c is None else str(c) for c in row]
                if any(cells):
                    lines.append("\t".join(cells).rstrip("\t"))
            if lines:
                pages.append(_native(f"Sheet: {ws.title}\n" + "\n".join(lines)))
                names.append(ws.title)
    finally:
        wb.close()
    return pages or [_native("")], {"sheets": ", ".join(names)}


def paginate_pptx(path):
    from pptx import Presentation
    try:
        prs = Presentation(str(path))
    except Exception as e:  # noqa: BLE001
        raise NonRetryable(f"Unreadable PPTX: {e}") from e
    pages = []
    for n, slide in enumerate(prs.slides, 1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                parts += [p.text for p in shape.text_frame.paragraphs if p.text.strip()]
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("Notes: " + notes)
        pages.append(_native(f"Slide {n}\n" + "\n".join(parts)))
    return pages or [_native("")], {"slides": str(len(pages))}


def paginate_csv(path):
    return [_native(path.read_text(encoding="utf-8-sig", errors="replace"))], {}


def paginate_json(path):
    try:
        data = json.loads(path.read_text(encoding="utf-8-sig"))
    except ValueError as e:
        raise NonRetryable(f"Invalid JSON: {e}") from e
    return [_native(json.dumps(data, indent=2, ensure_ascii=False))], {}


def paginate_xml(path):
    import defusedxml.ElementTree as ET
    try:
        root = ET.parse(str(path)).getroot()
    except Exception as e:  # noqa: BLE001 - includes defusedxml's entity-expansion rejections
        raise NonRetryable(f"Invalid or unsafe XML: {e}") from e
    lines: list[str] = []

    def walk(el, trail):
        here = f"{trail}/{el.tag.split('}')[-1]}"
        if el.text and el.text.strip():
            lines.append(f"{here}: {el.text.strip()}")
        for k, v in el.attrib.items():
            lines.append(f"{here}@{k}: {v}")
        for ch in el:
            walk(ch, here)

    walk(root, "")
    return [_native("\n".join(lines))], {}


def html_to_text(html: str) -> str:
    html = re.sub(r"(?is)<(script|style).*?</\1>", " ", html)
    return re.sub(r"[ \t]+", " ", re.sub(r"(?s)<[^>]+>", " ", html)).strip()


def paginate_eml(path):
    import email
    from email import policy
    try:
        with open(path, "rb") as f:
            msg = email.message_from_binary_file(f, policy=policy.default)
    except Exception as e:  # noqa: BLE001
        raise NonRetryable(f"Unreadable email: {e}") from e
    head = {k: str(msg[k]) for k in ("From", "To", "Cc", "Date", "Subject") if msg[k]}
    part = msg.get_body(preferencelist=("plain", "html"))
    body = ""
    if part is not None:
        body = part.get_content()
        if part.get_content_type() == "text/html":
            body = html_to_text(body)
    atts = [(a.get_filename() or "attachment", a.get_payload(decode=True) or b"") for a in msg.iter_attachments()]
    text = "\n".join(f"{k}: {v}" for k, v in head.items()) + "\n\n" + body
    return [_native(text)], {k.lower(): v for k, v in head.items()} | {"_attachments": atts}


def paginate_msg(path):
    import extract_msg
    try:
        m = extract_msg.Message(str(path))
    except Exception as e:  # noqa: BLE001
        raise NonRetryable(f"Unreadable MSG: {e}") from e
    try:
        head = {"From": m.sender, "To": m.to, "Cc": m.cc, "Date": str(m.date) if m.date else None, "Subject": m.subject}
        head = {k: v for k, v in head.items() if v}
        atts = [(a.longFilename or a.shortFilename or "attachment", a.data if isinstance(a.data, bytes) else b"")
                for a in m.attachments]
        text = "\n".join(f"{k}: {v}" for k, v in head.items()) + "\n\n" + (m.body or "")
    finally:
        m.close()
    return [_native(text)], {k.lower(): v for k, v in head.items()} | {"_attachments": atts}


PAGINATORS = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": paginate_xlsx,
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": paginate_pptx,
    "text/csv": paginate_csv, "application/json": paginate_json, "application/xml": paginate_xml,
    "message/rfc822": paginate_eml, "application/vnd.ms-outlook": paginate_msg,
}
