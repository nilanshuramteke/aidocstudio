import io
import json
from email.message import EmailMessage

import openpyxl
from pptx import Presentation

from adstudio.documents.service import Limits
from adstudio.ingestion.sniff import sniff

from .test_documents import make_png, make_zip, run_workers, settle, upload
from .test_queue import wait_for
from .test_search import drain, search, titles


def xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Expenses"
    ws.append(["Item", "Amount"])
    ws.append(["Laptop", 85000])
    ws2 = wb.create_sheet("Notes")
    ws2.append(["quarterly budget review"])
    b = io.BytesIO()
    wb.save(b)
    return b.getvalue()


def pptx_bytes() -> bytes:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Roadmap 2027"
    s.placeholders[1].text = "Launch the solar inverter line"
    s.notes_slide.notes_text_frame.text = "speaker note about tariffs"
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def eml_bytes(attach=True) -> bytes:
    m = EmailMessage()
    m["From"], m["To"], m["Subject"], m["Date"] = "a@x.com", "b@y.com", "Order confirmation 4471", "Mon, 2 Mar 2026 10:00:00 +0000"
    m.set_content("Your order for the espresso machine has shipped.")
    if attach:
        m.add_attachment(make_png("blue"), maintype="image", subtype="png", filename="receipt.png")
    return m.as_bytes()


def test_sniff_new_formats():
    assert sniff(xlsx_bytes()[:16]) == "zip"  # head alone cannot tell; needs the container check
    assert sniff(b"a,b,c\n1,2,3\n4,5,6\n") == "csv"
    assert sniff(b"name;qty\nx;1\ny;2\n") == "csv"
    assert sniff(b"just a sentence, with a comma\nand another line\n") == "txt"  # unequal delimiter counts
    assert sniff(b"From: a@x.com\nTo: b@y.com\nSubject: hi\n\nbody") == "eml"
    assert sniff(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\0" * 100) is None  # OLE that is not an Outlook message


def test_sniff_json_xml_with_files(tmp_path):
    j, x, bad = tmp_path / "a.bin", tmp_path / "b.bin", tmp_path / "c.bin"
    j.write_bytes(b'{"a": [1, 2]}')
    x.write_bytes(b"<?xml version='1.0'?><r><a>1</a></r>")
    bad.write_bytes(b"{not json")
    assert sniff(j.read_bytes(), full_path=j) == "json"
    assert sniff(x.read_bytes(), full_path=x) == "xml"
    assert sniff(bad.read_bytes(), full_path=bad) == "txt"


def test_office_and_data_formats_become_searchable(client):
    ids = {
        "xlsx": upload(client, "budget.xlsx", xlsx_bytes()),
        "pptx": upload(client, "deck.pptx", pptx_bytes()),
        "csv": upload(client, "people.csv", b"name,city\nAsha,Pune\nRavi,Chennai\n"),
        "json": upload(client, "cfg.json", json.dumps({"service": {"region": "ap-south-1", "retries": 3}}).encode()),
        "xml": upload(client, "feed.xml", b"<?xml version='1.0'?><feed><item>quarterly tariffs revision</item></feed>"),
    }
    assert all(r["status"] == "created" for r in ids.values()), ids
    run_workers(client)
    for r in ids.values():
        settle(client, r["document_id"])
    drain(client)
    mimes = {n: client.container.docs.get(r["document_id"])["mime"] for n, r in ids.items()}
    assert mimes["xlsx"].endswith("spreadsheetml.sheet") and mimes["csv"] == "text/csv" and mimes["xml"] == "application/xml"
    x = client.container.docs.get(ids["xlsx"]["document_id"])
    assert x["page_count"] == 2 and x["metadata"]["sheets"] == "Expenses, Notes"  # one page per sheet
    assert titles(search(client, "laptop")) == ["budget"] and titles(search(client, "quarterly budget")) == ["budget"]
    assert titles(search(client, "inverter")) == ["deck"] and titles(search(client, "tariffs", mode="keyword"))  # slide text + notes
    assert titles(search(client, "Chennai")) == ["people"] and titles(search(client, "ap-south-1")) == ["cfg"]
    assert client.container.docs.get(ids["pptx"]["document_id"])["page_count"] == 1


def test_email_becomes_doc_with_subject_title_and_attachment_children(client):
    raw = eml_bytes()  # boundaries are random per call, so reuse the exact bytes for the duplicate check
    r = upload(client, "msg.eml", raw)
    run_workers(client)
    settle(client, r["document_id"])
    drain(client)
    d = client.container.docs.get(r["document_id"])
    assert d["title"] == "Order confirmation 4471" and d["metadata"]["from"] == "a@x.com"
    assert titles(search(client, "espresso machine")) == ["Order confirmation 4471"]
    with client.container.db.read() as c:
        kids = c.execute("SELECT original_name, source, parent_id FROM documents WHERE parent_id=?", (r["document_id"],)).fetchall()
    assert [(k["original_name"], k["source"]) for k in kids] == [("receipt.png", "email")]
    assert upload(client, "again.eml", raw)["status"] == "duplicate"


def test_xml_entity_bombs_are_rejected_not_expanded(client):
    bomb = (b'<?xml version="1.0"?><!DOCTYPE lolz [<!ENTITY a "aaaaaaaaaa"><!ENTITY b "&a;&a;&a;&a;&a;&a;&a;&a;&a;&a;">'
            b'<!ENTITY c "&b;&b;&b;&b;&b;&b;&b;&b;&b;&b;">]><lolz>&c;</lolz>')
    r = upload(client, "bomb.xml", bomb)
    assert r["status"] == "created"  # sniffs as plain text (parse refused), so it is stored but never entity-expanded
    run_workers(client)
    settle(client, r["document_id"])
    assert client.container.docs.get(r["document_id"])["mime"] == "text/plain"


def test_zip_depth_and_shared_budget(client):
    lvl = make_zip({"leaf.png": make_png("green")})
    for _ in range(3):
        lvl = make_zip({"inner.zip": lvl})  # 3 wrappers: depth 0 -> 1 -> 2 -> leaf archive at depth 3
    r = upload(client, "deep.zip", lvl)

    def leaf(n):
        while n.get("children"):
            n = n["children"][0]
        return n

    assert leaf(r)["status"] == "rejected" and "deeper" in leaf(r)["reason"]  # past depth 3
    ok = upload(client, "ok.zip", make_zip({"a.zip": make_zip({"b.zip": make_zip({"c.png": make_png("red")})})}))
    assert leaf(ok)["status"] == "created"  # exactly 3 levels of nesting is allowed

    client.container.docs.limits = Limits(zip_max_entries=3)
    inner = make_zip({f"{i}.txt": f"doc number {i} here".encode() for i in range(3)})
    tree = upload(client, "tree.zip", make_zip({"x.zip": inner, "y.zip": inner, "z.txt": b"zzz"}))
    kids = {c["name"]: c for c in tree["children"]}
    assert kids["z.txt"]["status"] == "created"
    assert kids["x.zip"]["status"] == "rejected" and "too many files" in kids["x.zip"]["reason"]  # budget is shared
    assert kids["y.zip"]["status"] == "rejected"
