import csv
import io
import json
import os
import time
import zipfile
from pathlib import Path

import openpyxl
import pytest
from fastapi.testclient import TestClient

from adstudio.api.app import create_app
from adstudio.core import container as cmod
from adstudio.core.config import Config
from adstudio.core.errors import AppError
from adstudio.core.fakes import FakeLLM, FakeOCR
from adstudio.core.security import COOKIE_NAME
from adstudio.export.service import safe_cell
from adstudio.ops import backup as bk

from .test_documents import make_png, make_zip, run_workers, settle, upload
from .test_extraction import INVOICE_TEXT, pdf_with_lines
from .test_queue import wait_for
from .test_search import drain


def age(path: Path, seconds: float = 10) -> Path:
    t = time.time() - seconds
    os.utime(path, (t, t))
    return path


# ═══ watch folders ════════════════════════════════════════════
@pytest.fixture
def folder(tmp_path):
    f = tmp_path / "inbox_watch"
    f.mkdir()
    return f


def add(client, folder, **kw):
    r = client.post("/api/v1/watch-folders", json={"path": str(folder), **kw})
    assert r.status_code == 200, r.text
    return r.json()


def test_watch_folder_path_validation(client, folder):
    assert client.post("/api/v1/watch-folders", json={"path": "relative/dir"}).json()["code"] == "invalid_path"
    assert client.post("/api/v1/watch-folders", json={"path": str(folder / "missing")}).json()["code"] == "invalid_path"
    root = client.container.config.data_root
    assert client.post("/api/v1/watch-folders", json={"path": str(root / "files")}).json()["code"] == "invalid_path"
    assert client.post("/api/v1/watch-folders", json={"path": str(folder), "after_import": "explode"}).json()["code"] == "invalid_body"
    add(client, folder)
    assert client.post("/api/v1/watch-folders", json={"path": str(folder)}).status_code == 409


def test_scan_imports_new_files_once_and_skips_junk(client, folder):
    w = add(client, folder)
    age(folder.joinpath("a.png").write_bytes(make_png("red")) and folder / "a.png")
    age(folder.joinpath("b.txt").write_bytes(b"hello from the scanner folder") and folder / "b.txt")
    folder.joinpath("~$lock.docx").write_bytes(b"office lock file")
    folder.joinpath("half.pdf.part").write_bytes(b"%PDF-1.4 partial")
    folder.joinpath("junk.bin").write_bytes(b"\x00\x01\x02\x03binary")
    age(folder / "junk.bin")
    folder.joinpath("fresh.txt").write_bytes(b"still being written")  # just modified: must wait
    s = client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()
    assert (s["imported"], s["rejected"], s["skipped_recent"]) == (2, 1, 1)
    assert {d["title"] for d in client.get("/api/v1/documents").json()["items"]} == {"a", "b"}
    assert {d["source"] for d in [client.container.docs.get(i["id"]) for i in client.get("/api/v1/documents").json()["items"]]} == {"watch"}
    again = client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()
    assert again["imported"] == 0 and again["rejected"] == 0  # seen files (even rejected ones) are not retried
    age(folder / "fresh.txt")
    assert client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()["imported"] == 1  # now settled
    folder.joinpath("b.txt").write_bytes(b"hello from the scanner folder, edited")  # changed content is new content
    age(folder / "b.txt")
    assert client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()["imported"] == 1
    assert client.get("/api/v1/watch-folders").json()["items"][0]["last_scan"]["imported"] == 1


def test_after_import_move_and_delete_and_recursion(client, folder):
    sub = folder / "2026" / "march"
    sub.mkdir(parents=True)
    age(sub.joinpath("x.txt").write_bytes(b"deep file one") and sub / "x.txt")
    age(folder.joinpath("y.txt").write_bytes(b"top file two") and folder / "y.txt")
    w = add(client, folder, after_import="move")
    assert client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()["imported"] == 2
    assert not (sub / "x.txt").exists() and (folder / "_imported" / "2026" / "march" / "x.txt").exists()
    assert (folder / "_imported" / "y.txt").exists()
    age(folder.joinpath("y.txt").write_bytes(b"top file two") and folder / "y.txt")  # same content again, now a duplicate
    assert client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()["duplicate"] == 1
    assert (folder / "_imported" / "y (1).txt").exists()  # collision gets a suffix; _imported is not re-scanned

    flat = folder.parent / "flat"
    (flat / "inner").mkdir(parents=True)
    age(flat.joinpath("z.txt").write_bytes(b"flat file three") and flat / "z.txt")
    age((flat / "inner").joinpath("w.txt").write_bytes(b"inner file four") and flat / "inner" / "w.txt")
    d = add(client, flat, recursive=False, after_import="delete")
    assert client.post(f"/api/v1/watch-folders/{d['id']}/scan").json()["imported"] == 1
    assert not (flat / "z.txt").exists() and (flat / "inner" / "w.txt").exists()  # non-recursive left the subfolder alone
    assert client.patch(f"/api/v1/watch-folders/{d['id']}", json={"recursive": True}).json()["recursive"] is True
    assert client.delete(f"/api/v1/watch-folders/{d['id']}").status_code == 200


def test_unsupported_files_are_left_untouched_even_in_delete_mode(client, folder):
    age(folder.joinpath("junk.bin").write_bytes(b"\x00\x01\x02binary data") and folder / "junk.bin")
    w = add(client, folder, after_import="delete")
    assert client.post(f"/api/v1/watch-folders/{w['id']}/scan").json()["rejected"] == 1
    assert (folder / "junk.bin").exists()


def test_background_loop_picks_up_files(client, folder):
    add(client, folder)
    client.container.watch.start(0.2)
    try:
        age(folder.joinpath("late.txt").write_bytes(b"dropped while running") and folder / "late.txt")
        assert wait_for(lambda: client.get("/api/v1/documents").json()["items"], 10)
    finally:
        client.container.watch.stop()


# ═══ export ═══════════════════════════════════════════════════
def ready_invoices(client, n=2):
    ids = []
    for i in range(n):
        ids.append(upload(client, f"inv{i}.pdf", pdf_with_lines(INVOICE_TEXT.replace("INV-20491", f"INV-{i}")))["document_id"])
    run_workers(client)
    for d in ids:
        settle(client, d)
    drain(client)
    return ids


def test_csv_json_xlsx_export_contents(client):
    ids = ready_invoices(client)
    client.post("/api/v1/documents/bulk", json={"ids": ids[:1], "action": "tag", "params": {"tags": ["q1", "paid"]}})
    r = client.post("/api/v1/export", json={"document_ids": ids, "format": "csv"})
    assert r.status_code == 200 and r.headers["x-export-rows"] == "2" and r.headers["content-type"].startswith("text/csv")
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert [x["invoice_number"] for x in rows] == ["INV-0", "INV-1"] and rows[0]["total"] == "54280.00"
    assert rows[0]["tags"] == "paid, q1" and rows[0]["type"] == "Invoice" and rows[1]["tags"] == ""
    j = client.post("/api/v1/export", json={"document_ids": ids, "format": "json", "fields": ["total"]}).json()
    assert j[0]["fields"] == {"total": {"value": "54280.00", "confidence": j[0]["fields"]["total"]["confidence"], "status": "auto", "page": 1}}
    x = client.post("/api/v1/export", json={"document_ids": ids[:1], "format": "xlsx"})
    ws = openpyxl.load_workbook(io.BytesIO(x.content)).active
    header = [c.value for c in ws[1]]
    assert header[:3] == ["id", "title", "original_name"] and ws.cell(2, header.index("vendor") + 1).value == "ABC Private Limited"


def test_export_by_query_and_collection_and_errors(client):
    ids = ready_invoices(client)
    assert client.post("/api/v1/export", json={"query": {"q": "type:invoice"}, "format": "csv"}).headers["x-export-rows"] == "2"
    col = client.post("/api/v1/collections", json={"name": "C"}).json()
    client.post(f"/api/v1/collections/{col['id']}/documents", json={"document_ids": ids[:1]})
    assert client.post("/api/v1/export", json={"collection_id": col["id"], "format": "json"}).headers["x-export-rows"] == "1"
    assert client.post("/api/v1/export", json={"document_ids": ids, "format": "pdf"}).json()["code"] == "invalid_format"
    assert client.post("/api/v1/export", json={"document_ids": []}).json()["code"] == "nothing_to_export"
    assert client.post("/api/v1/export", json={}).json()["code"] == "invalid_body"
    assert client.post("/api/v1/documents/bulk", json={"ids": ids, "action": "export", "params": {"format": "csv"}}).json()["file"]["rows"] == 2


def test_formula_injection_is_neutralized(client):
    assert safe_cell("=HYPERLINK(\"http://x\")") == "'=HYPERLINK(\"http://x\")"
    assert safe_cell("@SUM(A1)") == "'@SUM(A1)" and safe_cell("-5.00") == "-5.00" and safe_cell("+91 98") == "'+91 98" and safe_cell(None) == ""
    ids = ready_invoices(client, 1)
    client.patch(f"/api/v1/documents/{ids[0]}/fields/customer", json={"value": '=cmd|" /C calc"!A0'})
    csv_text = client.post("/api/v1/export", json={"document_ids": ids, "format": "csv"}).content.decode("utf-8-sig")
    assert "'=cmd" in csv_text and ",=cmd" not in csv_text
    ws = openpyxl.load_workbook(io.BytesIO(client.post("/api/v1/export", json={"document_ids": ids, "format": "xlsx"}).content)).active
    cell = next(c for row in ws.iter_rows() for c in row if c.value and "cmd" in str(c.value))
    assert str(cell.value).startswith("'=") and cell.data_type != "f"  # stored as text, never as a formula


def test_export_download_blocks_path_traversal(client):
    ids = ready_invoices(client, 1)
    name = client.post("/api/v1/export", json={"document_ids": ids}).headers["content-disposition"].split("filename=")[-1].strip('"')
    assert client.get(f"/api/v1/exports/{name}").status_code == 200
    assert client.get("/api/v1/exports/..%2f..%2fdb%2fstudio.sqlite").status_code == 404


# ═══ backup / restore ═════════════════════════════════════════
def test_encryption_roundtrip_wrong_passphrase_and_truncation(tmp_path, monkeypatch):
    monkeypatch.setattr(bk, "CHUNK", 1000)  # force several chunks
    src = tmp_path / "plain.bin"
    src.write_bytes(os.urandom(5500))
    enc, dec = tmp_path / "x.adsbk", tmp_path / "out.bin"
    bk.encrypt_file(src, enc, "correct horse")
    assert enc.read_bytes()[:6] == b"ADSBK1" and src.read_bytes()[:50] not in enc.read_bytes()
    bk.decrypt_file(enc, dec, "correct horse")
    assert dec.read_bytes() == src.read_bytes()
    with pytest.raises(AppError) as e:
        bk.decrypt_file(enc, tmp_path / "bad.bin", "wrong")
    assert e.value.code == "bad_passphrase" and not (tmp_path / "bad.bin").exists()
    raw = enc.read_bytes()
    cut = tmp_path / "cut.adsbk"
    cut.write_bytes(raw[: len(raw) - 1020])  # drop the final chunk entirely
    with pytest.raises(AppError):
        bk.decrypt_file(cut, tmp_path / "cut.bin", "correct horse")
    flipped = bytearray(raw)
    flipped[100] ^= 1
    (tmp_path / "flip.adsbk").write_bytes(bytes(flipped))
    with pytest.raises(AppError):
        bk.decrypt_file(tmp_path / "flip.adsbk", tmp_path / "flip.bin", "correct horse")


def new_app(config, **kw):
    c = cmod.build_container(config, ocr=FakeOCR(), llm=FakeLLM(), start_workers=False, **kw)
    tc = TestClient(create_app(c), base_url="http://127.0.0.1")
    tc.__enter__()
    tc.cookies.set(COOKIE_NAME, c.auth.session_token)
    tc.container = c
    return tc


@pytest.mark.parametrize("passphrase", [None, "s3cret pass"])
def test_backup_and_restore_roundtrip(tmp_path, passphrase):
    cfg = Config.load(tmp_path / "data", workers=1, heartbeat_s=0.2, stale_after_s=0.5)
    a = new_app(cfg)
    first = upload(a, "keep.pdf", pdf_with_lines(INVOICE_TEXT))["document_id"]
    run_workers(a)
    settle(a, first)
    drain(a)
    b = a.post("/api/v1/backup", json={"passphrase": passphrase} if passphrase else {}).json()
    assert b["encrypted"] is bool(passphrase) and b["size"] > 0
    assert a.get("/api/v1/backups").json()["items"][0]["name"] == b["name"]
    second = upload(a, "later.txt", b"added after the backup was taken")["document_id"]
    settle(a, second)
    a.delete(f"/api/v1/documents/{first}")  # damage: delete one doc, add another
    assert a.post("/api/v1/restore", json={"name": "nope.zip"}).status_code == 404
    if passphrase:
        assert a.post("/api/v1/restore", json={"name": b["name"]}).json()["code"] == "passphrase_required"
        assert a.post("/api/v1/restore", json={"name": b["name"], "passphrase": "wrong"}).json()["code"] == "bad_passphrase"
    r = a.post("/api/v1/restore", json={"name": b["name"], "passphrase": passphrase}).json()
    assert r["restart_required"] and (cfg.data_root / bk.PENDING / "db" / "studio.sqlite").exists()
    a.container.stop()
    # "restart": building the container applies the staged restore before opening the DB
    c2 = new_app(cfg)
    assert c2.container.capabilities["restored_from_backup"] is True
    titles = {d["title"] for d in c2.get("/api/v1/documents").json()["items"]}
    assert titles == {"keep"}  # the later doc is gone, the deleted one is back
    assert c2.get(f"/api/v1/documents/{first}/file").status_code == 200  # originals restored too
    assert any(p.name.startswith("pre-restore-") for p in (cfg.data_root / "backups").iterdir())  # old state kept
    assert not (cfg.data_root / bk.PENDING).exists()
    c2.container.stop()


def test_restore_rejects_unsafe_or_broken_backups(tmp_path):
    root = tmp_path / "d"
    (root / "tmp").mkdir(parents=True)
    evil = root / "evil.zip"
    with zipfile.ZipFile(evil, "w") as z:
        z.writestr("manifest.json", "{}")
        z.writestr("db/studio.sqlite", b"x")
        z.writestr("../../escape.txt", b"nope")
    with pytest.raises(AppError) as e:
        bk.stage_restore(root, evil)
    assert e.value.code == "bad_backup" and not (tmp_path / "escape.txt").exists()
    nomani = root / "m.zip"
    with zipfile.ZipFile(nomani, "w") as z:
        z.writestr("hello.txt", "x")
    with pytest.raises(AppError):
        bk.stage_restore(root, nomani)
    corrupt = root / "c.zip"
    with zipfile.ZipFile(corrupt, "w") as z:
        z.writestr("manifest.json", "{}")
        z.writestr("db/studio.sqlite", b"this is not a sqlite database at all")
    with pytest.raises(AppError) as e2:
        bk.stage_restore(root, corrupt)
    assert e2.value.code == "bad_backup" and not (root / bk.PENDING).exists()  # a bad backup is never left staged
    garbage = root / "g.zip"
    garbage.write_bytes(b"not a zip")
    with pytest.raises(AppError):
        bk.stage_restore(root, garbage)
    assert bk.apply_pending_restore(root) is False


def test_scheduler_due_prune_and_job(client):
    c = client.container
    sched = c.scheduler
    assert sched.due() is False  # interval 0 = off
    client.patch("/api/v1/settings", json={"backup.interval_hours": 1.0, "backup.keep": 2})
    assert sched.due() is True  # none yet
    assert sched.tick() is True
    run_workers(client)
    assert wait_for(lambda: any(b["auto"] for b in bk.list_backups(c.config.data_root)), 20)
    assert sched.due() is False  # a fresh auto backup exists
    for _ in range(3):  # more autos; the handler prunes down to `keep`
        c.queue.enqueue("backup", {"auto": True})
        time.sleep(1.1)  # distinct timestamps
    drain(client)
    assert len([b for b in bk.list_backups(c.config.data_root) if b["auto"]]) == 2
    manual = bk.create_backup(c.config.data_root, c.db)
    assert manual.exists() and manual.name in {b["name"] for b in bk.list_backups(c.config.data_root)}  # manual ones are never pruned
    assert client.get(f"/api/v1/backups/{manual.name}").status_code == 200
    assert client.get("/api/v1/backups/..%2fconfig.toml").status_code == 404
    assert client.get("/api/v1/audit").json()["items"] is not None
